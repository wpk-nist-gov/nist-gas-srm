"""standard analysis models"""
# pylint: disable=abstract-method

from functools import partial
from typing import TYPE_CHECKING, Annotated, ClassVar, TypeAlias, cast

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string
from pydantic import BeforeValidator
from sqlmodel import (
    Field,
    SQLModel,
)
from sqlmodel._compat import SQLModelConfig  # ruff:ignore[import-private-name]

from nist_gas_srm.core.excel_interface import SheetNames, SQLDataFrameInterface
from nist_gas_srm.core.excel_utils import (
    get_fill_from_cell,
    get_value_from_worksheet,
    simple_write_to_excel,
    skipper,
)
from nist_gas_srm.core.standard_analysis import FIT_NAME_DEGREE_MAPPING
from nist_gas_srm.core.typing_compat import override
from nist_gas_srm.core.validate import (
    validate_nan_to_none,
)

if TYPE_CHECKING:
    from openpyxl.cell.cell import Cell

from .keys import (
    IDPrimaryKeyPublic,
    SRMDataForeignKey,
    SRMDataForeignKeyUpdate,
    StandardAnalysisForeignKey,
)
from .utils import (
    to_pascal,
)


# * StandardAnalysis ----------------------------------------------------------
# * root table
class _StandardAnalysisForeignKeyUpdate(SQLModel):
    pass


# NOTE(wpk): include SQLModel here to make linter happy
class StandardAnalysisBase(SRMDataForeignKey, SQLModel):
    """Points back to srmdata"""


class StandardAnalysisPublic(StandardAnalysisBase, IDPrimaryKeyPublic):
    pass


class StandardAnalysisCreate(StandardAnalysisBase):
    pass


class StandardAnalysisUpdate(SRMDataForeignKeyUpdate):
    pass


# ** Subtables
class ParamsBase(StandardAnalysisForeignKey):
    model_config = SQLModelConfig(populate_by_name=True)
    degree: int | None

    assumed_lot_standard: float | None
    assumed_lot_standard_uncert: float | None

    montecarlo_estimate: float | None
    montecarlo_estimate_uncert: float | None
    montecarlo_slope: float | None
    montecarlo_slope_uncert: float | None
    montecarlo_intercept: float | None
    montecarlo_intercept_uncert: float | None


class ParamsPublic(ParamsBase, IDPrimaryKeyPublic):
    pass


class ParamsCreate(ParamsBase, SQLDataFrameInterface):
    dataframe_name = "standard_analysis.params"
    sheet_name = SheetNames.standard_analysis

    # NOTE row and column are base 1, NOT base 0
    row_column_mapping: ClassVar[dict[str, tuple[int, int]]] = {
        "degree": (4, 12),
        "assumed_lot_standard": (1, 12),
        "assumed_lot_standard_uncert": (1, 13),
        "montecarlo_estimate": (8, 12),
        "montecarlo_estimate_uncert": (8, 13),
        "montecarlo_slope": (9, 12),
        "montecarlo_slope_uncert": (9, 13),
        "montecarlo_intercept": (10, 12),
        "montecarlo_intercept_uncert": (10, 13),
    }

    @override
    @classmethod
    def excel_to_dataframe(cls, excelfile: pd.ExcelFile) -> pd.DataFrame | None:

        get_value = partial(
            get_value_from_worksheet, xls=excelfile, sheet_name=cls.sheet_name
        )

        data = {
            k: get_value(rowx=rowx - 1, colx=colx - 1)
            for k, (rowx, colx) in cls.row_column_mapping.items()
        }

        return pd.DataFrame([data])

    @override
    @classmethod
    def dataframe_to_excel(cls, obj: pd.DataFrame, workbook: Workbook) -> None:

        if len(obj) != 1:
            msg = "should be length one dataframe"
            raise ValueError(msg)
        data = obj.to_dict(orient="records")[0]

        worksheet = workbook[cls.sheet_name]

        for key, (rowx, colx) in cls.row_column_mapping.items():
            if key in data:
                target_cell = cast("Cell", worksheet.cell(row=rowx, column=colx))
                target_cell.value = data[key]


class ParamsUpdate(_StandardAnalysisForeignKeyUpdate):
    pass


# ** GenLineParams subtable of StandardAnalysis
class GenLineParamsBase(StandardAnalysisForeignKey):
    model_config = SQLModelConfig(
        alias_generator=to_pascal,
        populate_by_name=True,
    )

    name: str
    value: Annotated[float | None, BeforeValidator(validate_nan_to_none)]
    stderr: float


class GenLineParamsPublic(
    GenLineParamsBase,
    IDPrimaryKeyPublic,
):
    pass


class GenLineParamsCreate(GenLineParamsBase, SQLDataFrameInterface):
    dataframe_name = "standard_analysis.genline_params"
    sheet_name = SheetNames.standard_analysis

    @override
    @classmethod
    def excel_to_dataframe(
        cls, excelfile: pd.ExcelFile, degree: int | None = None
    ) -> pd.DataFrame | None:
        # get type of fit
        if degree is None:
            fit_name = get_value_from_worksheet(
                excelfile, cls.sheet_name, rowx=21, colx="T"
            )

            degree = FIT_NAME_DEGREE_MAPPING[fit_name]

        nrows = int((degree + 1) * (1 + degree / 2) + 1)

        return cls._get_optional_frame(
            excelfile,
            usecols="T:V",
            header=None,
            skiprows=skipper(include=range(23, 23 + nrows)),
            names=["name", "value", "stderr"],
        )

    @override
    @classmethod
    def dataframe_to_excel(cls, obj: pd.DataFrame, workbook: Workbook) -> None:

        worksheet = workbook[cls.sheet_name]

        simple_write_to_excel(
            obj,
            worksheet=workbook[cls.sheet_name],
            start=(24, "T"),
            fill_from=(23, "T"),
            header=False,
        )

        degree = int((-1 + np.sqrt(1 - 4 * 2 * (1 - len(obj)))) / 2 - 1)
        fit_name = {v: k for k, v in FIT_NAME_DEGREE_MAPPING.items()}[degree]

        fill = get_fill_from_cell(worksheet.cell(row=23, column=20))
        cell = cast(
            "Cell", worksheet.cell(row=22, column=column_index_from_string("T"))
        )
        cell.value = fit_name
        cell.fill = fill


class GenLineParamsUpdate(_StandardAnalysisForeignKeyUpdate):
    name: str | None = None
    value: float | None = None
    stderr: float | None = None


# ** GenLineSolution subtable of StandardAnalysis
class GenLineSolutionBase(StandardAnalysisForeignKey):
    model_config = SQLModelConfig(
        alias_generator=to_pascal,
        populate_by_name=True,
    )

    x_solution: float = Field(validation_alias="X-Solution")
    y_solution: float = Field(validation_alias="Y-Solution")


class GenLineSolutionPublic(
    GenLineSolutionBase,
    IDPrimaryKeyPublic,
):
    pass


class GenLineSolutionCreate(GenLineSolutionBase, SQLDataFrameInterface):
    dataframe_name = "standard_analysis.genline_solution"
    sheet_name = SheetNames.standard_analysis

    @override
    @classmethod
    def excel_to_dataframe(
        cls, excelfile: pd.ExcelFile, degree: int | None = None
    ) -> pd.DataFrame | None:

        if degree is None:
            fit_name = get_value_from_worksheet(
                excelfile, cls.sheet_name, rowx=21, colx="T"
            )
            degree = FIT_NAME_DEGREE_MAPPING[fit_name]

        nrows = int((degree + 1) * (1 + degree / 2) + 1)
        start = 24 - 1 + nrows + 2 - 1

        out = cls._get_optional_frame(
            excelfile, usecols="T:X", skiprows=skipper(include=range(start, start + 10))
        )
        if out is None:
            return out

        return out.dropna()[["X-Solution", "Y-Solution"]]

    @override
    @classmethod
    def dataframe_to_excel(cls, obj: pd.DataFrame, workbook: Workbook) -> None:

        worksheet = workbook[cls.sheet_name]
        fit_name = str(worksheet.cell(row=22, column=20).value)
        degree = FIT_NAME_DEGREE_MAPPING[fit_name]

        nrows = int((degree + 1) * (1 + degree / 2) + 1)
        start = 24 + nrows + 2 - 1

        def _get_pass_formula(i: int, offset: int, start: int) -> str:
            j = start + 1 + i - offset
            return f'=IF(AND(ABS(W{j}-U{j})<2 * W{i}, ABS(V{j}-T{j})<2*U{i}), "PASS", "FAIL")'

        df = obj.assign(
            X=[f"=T{i}" for i in range(2, 2 + len(obj))],
            Y=[f"=V{i}" for i in range(2, 2 + len(obj))],
            uTest=[
                _get_pass_formula(i, offset=2, start=start)
                for i in range(2, 2 + len(obj))
            ],
        )[["X", "Y", "X-Solution", "Y-Solution", "uTest"]]

        simple_write_to_excel(
            df,
            worksheet=worksheet,
            start=(start, "T"),
            fill_from=(23, "T"),
        )


class GenLineSolutionUpdate(_StandardAnalysisForeignKeyUpdate):
    x_solution: float | None = None
    y_solution: float | None = None


# ** GenLine Eval Base
# NOTE: I'd really rather calculate this...
class GenLineEvalBase(StandardAnalysisForeignKey):
    model_config = SQLModelConfig(
        alias_generator=to_pascal,
        populate_by_name=True,
    )

    y_eval: float = Field(validation_alias="Yeval")
    y_eval_uncert: float = Field(validation_alias="uYeval")


class GenLineEvalPublic(
    GenLineEvalBase,
    IDPrimaryKeyPublic,
):
    pass


class GenLineEvalCreate(GenLineEvalBase, SQLDataFrameInterface):
    dataframe_name = "standard_analysis.genline_eval"
    sheet_name = SheetNames.standard_analysis

    @override
    @classmethod
    def excel_to_dataframe(
        cls, excelfile: pd.ExcelFile, degree: int | None = None
    ) -> pd.DataFrame | None:

        if degree is None:
            fit_name = get_value_from_worksheet(
                excelfile, cls.sheet_name, rowx=21, colx="T"
            )
            degree = FIT_NAME_DEGREE_MAPPING[fit_name]

        nrows = int((degree + 1) * (1 + degree / 2) + 1)
        start = 24 - 1 + nrows + 2 - 1

        out = cls._get_optional_frame(
            excelfile, usecols="T:X", skiprows=skipper(include=range(start, start + 30))
        )
        if out is None:
            return out

        query = out.query("X=='Xin'")

        if len(query) != 1:
            msg = f"DataFrame length {len(query)} != 1."
            raise ValueError(msg)

        df = out.loc[query.index[0] :].iloc[1:]
        df.columns = query.iloc[0].to_numpy()

        return df[["Yeval", "uYeval"]]

    @override
    @classmethod
    def dataframe_to_excel(cls, obj: pd.DataFrame, workbook: Workbook) -> None:

        worksheet = workbook[cls.sheet_name]
        fit_name = str(worksheet.cell(row=22, column=20).value)
        degree = FIT_NAME_DEGREE_MAPPING[fit_name]

        nrows = int((degree + 1) * (1 + degree / 2) + 1)
        start = 24 + nrows + 2 - 1

        def find_end_of_solution_table() -> int:
            get_cell = partial(worksheet.cell, column=22)

            if get_cell(row=start).value != "X-Solution":
                msg = "Must write GenLineSolution before GenLineEval"
                raise ValueError(msg)

            for row in range(start + 1, start + 30):
                cell = get_cell(row)
                if cell.value is None:
                    return row

            msg = "Failed to find end of data."
            raise ValueError(msg)

        new_start = find_end_of_solution_table() + 1

        df = obj.assign(
            Xin=[f"=T{i}" for i in range(12, 12 + len(obj))],
            uXin=[f"=U{i}" for i in range(12, 12 + len(obj))],
        )[["Xin", "uXin", "Yeval", "uYeval"]]

        simple_write_to_excel(
            df,
            worksheet=worksheet,
            start=(new_start, "T"),
            fill_from=(23, "T"),
        )

        for col, target_col in [
            (12, "V"),
            (13, "W"),
        ]:
            cell = cast("Cell", worksheet.cell(row=5, column=col))
            cell.value = f"={target_col}{new_start + 1}"


class GenLineEvalUpdate(_StandardAnalysisForeignKeyUpdate):
    y_eval: float | None = None
    y_eval_uncert: float | None = None


StandardAnalysisSubTableCreate: TypeAlias = (
    ParamsCreate | GenLineParamsCreate | GenLineSolutionCreate | GenLineEvalCreate
)


# * Complete
class CompletePublic(StandardAnalysisPublic):
    params: list[ParamsPublic] = []
    genline_params: list[GenLineParamsPublic] = []
    genline_solution: list[GenLineSolutionPublic] = []
    genline_eval: list[GenLineEvalPublic] = []


class CompleteCreate(StandardAnalysisCreate):
    params: list[ParamsCreate] = []
    genline_params: list[GenLineParamsCreate] = []
    genline_solution: list[GenLineSolutionCreate] = []
    genline_eval: list[GenLineEvalCreate] = []
